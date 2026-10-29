"""Read supported uploads and keep the active chat files together."""
import base64
import hashlib
import warnings
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import pandas as pd
from PIL import Image, ImageOps
from app.rag.documents import DocumentError
from app.spreadsheet.analysis import read_table, detect_header_row

TEXT = {'.txt', '.md', '.rst', '.log', '.json', '.jsonl', '.yaml', '.yml',
        '.xml', '.html', '.htm', '.py', '.js', '.ts', '.css', '.sql', '.ini'}
IMAGES = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.gif', '.tif', '.tiff',
          '.ico', '.avif', '.heic', '.heif'}
TABLES = {'.xlsx', '.xls', '.xlsb', '.ods', '.csv', '.tsv'}
MAX_BYTES = 10 * 1024 * 1024


def check_archive(data):
    with ZipFile(BytesIO(data)) as archive:
        if len(archive.infolist()) > 5000 or sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
            raise DocumentError('Expanded document exceeds the 50 MB / 5,000 entry limit.')


def image_url(data):
    # Convert images to a supported format and use the first frame.
    import pillow_heif
    pillow_heif.register_heif_opener()
    with warnings.catch_warnings():
        warnings.simplefilter('error', Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as image:
            if image.width * image.height > 20_000_000:
                raise DocumentError('Image exceeds 20 megapixels. Resize it first.')
            image.seek(0)
            normalized = ImageOps.exif_transpose(image).convert('RGB')
            normalized.thumbnail((1600, 1600))
            output = BytesIO()
            normalized.save(output, format='JPEG', quality=85)
    return 'data:image/jpeg;base64,' + base64.b64encode(output.getvalue()).decode('ascii')


class AttachmentStore:
    """Keep active attachments in memory; the optional RAG backend also stores its source files on disk."""
    def __init__(self, documents):
        self.schema_version = 2
        self.documents = documents
        self.items = {}
        self.tables = {}
        self.images = {}
        self.image_answers = {}
        self.image_notes = {}
        self.vision_state = {'cooldown': {}, 'events': []}

    def clear(self):
        self.items.clear()
        self.tables.clear()
        self.images.clear()
        self.image_answers.clear()
        self.image_notes.clear()
        self.vision_state = {'cooldown': {}, 'events': []}
        self.documents.clear()

    def cached_image_answer(self, filename, question):
        key=(filename, self.items.get(filename, {}).get('digest'), ' '.join(question.lower().split()))
        return self.image_answers.get(key)

    def remember_image_answer(self, filename, question, answer):
        key=(filename, self.items.get(filename, {}).get('digest'), ' '.join(question.lower().split()))
        if len(self.image_answers)>=30:
            self.image_answers.pop(next(iter(self.image_answers)))
        self.image_answers[key]=answer
        self.image_notes[filename]=answer[:1200]

    def manifest(self):
        return list(self.items.values())

    def add(self, filename, data):
        name = filename.replace('\\', '/').split('/')[-1]
        suffix = Path(name).suffix.lower()
        if not data or len(data) > MAX_BYTES:
            raise DocumentError('Use a non-empty file up to 10 MB.')
        if name not in self.items and len(self.items) >= 10:
            raise DocumentError('Use at most 10 attachments per chat. Clear attachments to add more.')
        digest = hashlib.sha256(data).hexdigest()
        if self.items.get(name, {}).get('digest') == digest:
            state = getattr(self.documents, 'backend_status', {}).get(name, {})
            if state and state.get('retrieval') != 'RAG-Anything':
                # Reattaching is an explicit retry, not a new copy of the same file.
                if self.items[name].get('kind') == 'document' and hasattr(self.documents, 'add_file'):
                    self.documents.add_file(name, data)
                elif name in self.tables and hasattr(self.documents, 'add_table_context'):
                    self.documents.add_table_context(name, self.tables[name])
                self.items[name]['knowledge_backend'] = dict(self.documents.backend_status.get(name, {}))
            return self.items[name]
        if name in self.items:
            raise DocumentError('A different file with this name is loaded. Rename it or clear attachments first.')
        details = {}
        try:
            if suffix == '.pdf':
                self.documents.add(name, data)
                kind = 'document'
                details = {'location': 'PDF page', 'reader': 'pypdf'}
            elif suffix in TABLES:
                if suffix in {'.xlsx', '.ods', '.xlsb'}:
                    check_archive(data)
                if suffix in {'.csv', '.xlsx'}:
                    header = detect_header_row(data, name)
                    frame = read_table(data, name, header_row=header)
                    sheet_name = 'CSV'
                    if suffix == '.xlsx':
                        with pd.ExcelFile(BytesIO(data)) as book:
                            sheet_name = book.sheet_names[0]
                            sheet_count = len(book.sheet_names)
                    else:
                        sheet_count = 1
                else:
                    header = 1
                    if suffix == '.tsv':
                        frame = pd.read_csv(BytesIO(data), sep='\t', nrows=50001)
                        sheet_name, sheet_count = 'TSV', 1
                    else:
                        with pd.ExcelFile(BytesIO(data)) as book:
                            sheet_name, sheet_count = book.sheet_names[0], len(book.sheet_names)
                            frame = pd.read_excel(book, sheet_name=0, nrows=50001)
                    frame = frame.dropna(axis=1, how='all').dropna(axis=0, how='all')
                if frame.empty or len(frame) > 50000 or len(frame.columns) > 100:
                    raise DocumentError('Table limit: 50,000 rows and 100 columns.')
                frame.columns = frame.columns.map(str)
                if frame.columns.duplicated().any():
                    raise DocumentError('Table column names must be unique.')
                self.tables[name] = frame
                kind = 'table'
                details = {'rows': len(frame), 'columns': list(frame.columns), 'header_row': header,
                           'sheet': sheet_name, 'coverage': 'First worksheet only', 'sheet_count': sheet_count,
                           'reader': 'pandas (local calculations; cached Excel formula values only)'}
            elif suffix in IMAGES:
                self.images[name] = image_url(data)
                kind = 'image'
                details = {'reader': 'Pillow + configured vision service', 'coverage': 'First frame, resized to max 1600 pixels'}
            else:
                sections = []
                location = 'section'
                if suffix == '.docx':
                    check_archive(data)
                    from docx import Document
                    document = Document(BytesIO(data))
                    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
                    paragraphs.extend(' | '.join(c.text for c in row.cells)
                                      for table in document.tables for row in table.rows)
                    sections = ['\n'.join(paragraphs)]
                elif suffix == '.pptx':
                    check_archive(data)
                    from pptx import Presentation
                    slides = Presentation(BytesIO(data)).slides
                    if len(slides) > 100:
                        raise DocumentError('Use at most 100 slides.')
                    sections = ['\n'.join(shape.text if shape.has_text_frame else
                                '\n'.join(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                                if shape.has_table else '' for shape in slide.shapes) for slide in slides]
                    location = 'slide'
                elif suffix in TEXT:
                    sections = [data.decode('utf-8-sig')]
                else:
                    raise DocumentError('Unsupported format. Use PDF, DOCX, PPTX, XLSX/XLS/XLSB/ODS, CSV/TSV, text/code or common images. Legacy DOC/PPT, SVG/RAW/PSD, audio/video, archives and executables are not supported.')
                if suffix in {'.docx','.pptx'} and hasattr(self.documents,'add_file'):
                    try:self.documents.add_file(name,data)
                    except DocumentError:self.documents.add_sections(name,sections)
                else:self.documents.add_sections(name,sections)
                kind = 'document'
                details = {'location': location, 'reader': 'Local text extraction',
                           'coverage': 'Text and tables only; embedded pictures are not analyzed'}
        except DocumentError:
            raise
        except Exception:
            raise DocumentError('Could not read this file. It may be damaged, encrypted, non-UTF-8 or require conversion. No content was analyzed.') from None
        if kind=='image' and hasattr(self.documents,'add_file'):
            try:self.documents.add_file(name,data)
            except DocumentError:pass  # The existing direct-vision route remains usable.
        if kind=='table' and hasattr(self.documents,'add_table_context'):
            self.documents.add_table_context(name,self.tables[name])
        status=getattr(self.documents,'backend_status',{}).get(name)
        if status:details['knowledge_backend']=dict(status)
        item = {'name': name, 'kind': kind, 'digest': digest, **details}
        self.items[name] = item
        return item


def upgrade_attachment_store(store):
    """Keep older callers working; the UI uses the migration module."""
    from app.attachments.migration import upgrade_attachment_store as upgrade
    return upgrade(store)
