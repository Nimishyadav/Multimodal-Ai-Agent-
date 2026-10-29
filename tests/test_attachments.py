import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from PIL import Image
from docx import Document
from pptx import Presentation
from langchain_core.messages import AIMessage
from app.attachments.store import AttachmentStore
from app.attachments.vision import describe_image
from app.rag.documents import DocumentIndex, DocumentError
from app.agent.session import Conversation
from app.history.store import HistoryStore
from app.tools.document_search import create_document_tool
from app.multi_agent.system import handle_multi_agent
from test_multi_agent import Model, call
from test_workflows import FakeSearch
from test_chat_router import PriceStub


class AttachmentTests(unittest.TestCase):
    def store(self): return AttachmentStore(DocumentIndex())

    def test_word_and_slides_retrieve_actual_text(self):
        store = self.store()
        document = Document()
        document.add_paragraph('Solar panels convert sunlight into electricity.')
        document.add_table(rows=1, cols=1).cell(0, 0).text = 'Efficiency 20 percent'
        data = BytesIO(); document.save(data)
        store.add('guide.docx', data.getvalue())
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'Green hydrogen'
        data = BytesIO(); deck.save(data)
        store.add('slides.pptx', data.getvalue())
        self.assertIn('Solar', store.documents.search('sunlight')['passages'][0]['text'])
        self.assertEqual(store.documents.search('hydrogen')['passages'][0]['page'], 1)

    def test_csv_tsv_and_multiple_table_selection(self):
        store = self.store()
        store.add('one.csv', b'Tax\n10\n20\n')
        store.add('two.tsv', b'Tax\tName\n7\tAlice\n8\tBob\n')
        self.assertEqual(store.tables['two.tsv']['Tax'].sum(), 15)
        model = Model(call('delegate', {'agent':'data','task':'Sum tax in two.tsv'}, 'h'),
            call('inspect_table', {}, 'i'),
            call('calculate_table', {'table_name':'two.tsv','plan_json':json.dumps({'operation':'sum','column':'Tax'})}, 'c'),
            AIMessage(content='Invented 999'), AIMessage(content='Invented 999'))
        with tempfile.TemporaryDirectory() as folder:
            conv=Conversation(model, history_store=HistoryStore(Path(folder)/'history.sqlite3'))
            try:
                answer,_=handle_multi_agent('Excel total tax for two.tsv',conv,model,FakeSearch(),PriceStub(),attachments=store)
                self.assertIn('15', answer)
                self.assertIn('two.tsv', answer)
                self.assertNotIn('999', answer)
            finally: conv.close()

    def test_common_images_normalize_without_api_call(self):
        store=self.store()
        for ext,fmt in [('png','PNG'),('jpg','JPEG'),('webp','WEBP'),('gif','GIF'),('tiff','TIFF'),('bmp','BMP')]:
            data=BytesIO(); Image.new('RGB',(20,20),'red').save(data,format=fmt)
            item=store.add('image.'+ext,data.getvalue())
            self.assertEqual(item['kind'],'image')
            self.assertTrue(store.images['image.'+ext].startswith('data:image/jpeg;base64,'))

    def test_invalid_duplicate_and_limits(self):
        store=self.store()
        for name,data in [('file.exe',b'MZ'),('bad.png',b'not an image'),('blank.txt',b''),('bad.docx',b'not a zip')]:
            with self.assertRaises(DocumentError): store.add(name,data)
        store.add('notes.txt',b'Original text')
        store.add('notes.txt',b'Original text')
        with self.assertRaises(DocumentError): store.add('notes.txt',b'Different text')
        self.assertEqual(len(store.manifest()),1)
        with self.assertRaises(DocumentError): store.add('huge.txt',b'a'*(10*1024*1024+1))

    def test_vision_sends_actual_image_and_question(self):
        class Vision:
            def invoke(self,messages):
                self.messages=messages
                return AIMessage(content='A red square.')
        model=Vision()
        result=describe_image('data:image/jpeg;base64,test','What color?',model)
        self.assertEqual(result,'A red square.')
        self.assertEqual(model.messages[-1].content[1]['image_url']['url'],'data:image/jpeg;base64,test')

    def test_image_delegation_keeps_vision_evidence(self):
        store=self.store()
        data=BytesIO(); Image.new('RGB',(10,10),'red').save(data,format='PNG')
        store.add('red.png',data.getvalue())
        model=Model(call('delegate',{'agent':'images','task':'Describe red.png'},'h'),
            call('analyze_image',{'filename':'red.png','question':'What color?'},'v'),
            AIMessage(content='Fabricated blue square'),AIMessage(content='Fabricated blue square'))
        vision=Model(AIMessage(content='A red square.'))
        with tempfile.TemporaryDirectory() as folder:
            conv=Conversation(model,history_store=HistoryStore(Path(folder)/'history.sqlite3'))
            try:
                answer,_=handle_multi_agent('Describe this image',conv,model,FakeSearch(),PriceStub(),attachments=store,vision_model=vision)
                self.assertIn('red square',answer)
                self.assertNotIn('blue square',answer)
                self.assertIn('Image Agent selected tool: analyze_image',conv.activity)
            finally:conv.close()

    def test_word_source_uses_section_and_exports(self):
        store=self.store(); store.add('notes.txt',b'Solar panels convert sunlight into electricity.')
        model=Model(call('delegate',{'agent':'documents','task':'Explain document sunlight'},'h'),
            call('search_documents',{'query':'sunlight'},'s'), AIMessage(content='Fabricated'), AIMessage(content='Fabricated'))
        with tempfile.TemporaryDirectory() as folder:
            history=HistoryStore(Path(folder)/'history.sqlite3')
            conv=Conversation(model,history_store=history,document_index=store.documents,
                              additional_tools=[create_document_tool(store.documents)])
            try:
                answer,_=handle_multi_agent('Explain document sunlight',conv,model,FakeSearch(),PriceStub(),attachments=store)
                self.assertIn('section 1',answer)
                self.assertNotIn('PDF page',answer)
                self.assertIn('notes.txt',history.export_markdown(conv.user_id,conv.thread_id))
            finally:conv.close()
