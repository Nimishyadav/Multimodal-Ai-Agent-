import asyncio
from io import BytesIO
from pathlib import Path
import tempfile
import types
import unittest
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from app.attachments.store import AttachmentStore
from app.rag.documents import DocumentIndex
from app.rag.anything_rag import AnythingRAGIndex
from app.rag import anything_rag_worker as worker

BODY=('Document content for large attachment verification. '*260)
MARKER='violetkiwifinalmarker'

@patch.dict('os.environ', {'RAG_EMBEDDING_PROVIDER':'local','RAG_PROCESSING_MODE':'local'})
class LargeDocumentTests(unittest.TestCase):
 def test_twenty_page_word_exceeds_old_limit_and_keeps_end(self):
  doc=Document()
  for i in range(20):
   doc.add_paragraph(BODY)
   if i<19:doc.add_page_break()
  doc.add_paragraph(MARKER)
  data=BytesIO();doc.save(data)
  store=AttachmentStore(DocumentIndex());store.add('large.docx',data.getvalue())
  result=store.documents.search(MARKER)
  self.assertEqual(result['status'],'ok')
  self.assertTrue(any(MARKER in item['text'] for item in result['passages']))
  self.assertGreater(store.documents.documents()[0]['chunks'],200)

 def test_twenty_page_pdf_exceeds_old_limit_and_keeps_last_page(self):
  writer=PdfWriter()
  font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
  for i in range(20):
   page=writer.add_blank_page(width=612,height=842)
   page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
   lines=['Large PDF extraction verification and source evidence. '*3]*80
   if i==19:lines.append(MARKER)
   commands=['BT /F1 6 Tf 10 820 Td']
   for line in lines:commands.append('('+line+') Tj 0 -9 Td')
   commands.append('ET');stream=DecodedStreamObject();stream.set_data('\n'.join(commands).encode())
   page[NameObject('/Contents')]=writer._add_object(stream)
  data=BytesIO();writer.write(data)
  index=DocumentIndex();index.add('large.pdf',data.getvalue())
  self.assertEqual(index.documents()[0]['text_pages'],20)
  result=index.search(MARKER)
  self.assertEqual(result['passages'][0]['page'],20)
  self.assertIn(MARKER,result['passages'][0]['text'])

 def test_anything_rag_accepts_large_parsed_evidence(self):
  index=AnythingRAGIndex()
  index._install_pages('large.pdf','testhash',[{'page':i+1,'text':BODY+(MARKER if i==19 else '')} for i in range(20)])
  self.assertEqual(index.search(MARKER)['passages'][0]['page'],20)

 def test_optional_worker_accepts_large_parsed_content(self):
  class Engine:
   lightrag = None
   async def _ensure_lightrag_initialized(self):
    return {'success': True}
   async def parse_document(self,*args,**kwargs):
    return [{'type':'text','page_idx':i,'text':BODY} for i in range(20)], 'document-id'
  from unittest.mock import AsyncMock
  with tempfile.TemporaryDirectory() as temporary:
   root=Path(temporary);folder=root/'data'/'multimodal_rag'/'testjob'
   (folder/'source').mkdir(parents=True);(folder/'source'/'input.pdf').write_bytes(b'fixture')
   with patch.object(worker,'ROOT',root),patch('app.rag.anything_rag_models.build_engine',new=AsyncMock(return_value=Engine())):
    result=asyncio.run(worker.execute({'action':'parse','folder':str(folder),'name':'large.pdf'}))
   self.assertTrue(result['ok'])
   self.assertEqual(len(result['pages']),20)
   self.assertGreater(sum(len(p['text']) for p in result['pages']),200000)

if __name__=='__main__':unittest.main()
